
from flask import Flask, render_template, request, send_file
from docx import Document
from io import BytesIO
import requests
import json
import logging

app = Flask(__name__)
logging.basicConfig(level=logging.INFO)

OLLAMA_URL = "http://127.0.0.1:11434/api/generate"
MODEL_NAME = "gemma2:2b"
OLLAMA_TIMEOUT = 300


def validate_proposal(data, topic="the selected research topic"):
    """Validate and repair missing or incorrectly formatted sections."""

    if not isinstance(data, dict):
        raise ValueError("The AI response is not a JSON object.")

    defaults = {
        "title": f"A Research Study on {topic}",
        "summary": (
            f"This proposal investigates {topic} and its important "
            "aspects using a structured research approach."
        ),
        "problem": (
            f"This study examines key challenges related to {topic} "
            "and identifies areas that require further investigation."
        ),
        "objectives": [
            f"Understand the main aspects of {topic}.",
            f"Analyze challenges related to {topic}.",
            f"Identify possible improvements related to {topic}."
        ],
        "methodology": (
            "1. Review relevant literature. "
            "2. Collect and analyze appropriate information. "
            "3. Interpret the findings in relation to the objectives."
        ),
        "expected_outcomes": (
            f"The study aims to improve understanding of {topic} "
            "and identify possible areas for future research."
        ),
        "research_ideas": [
            f"Investigate current challenges in {topic}.",
            f"Compare existing approaches to {topic}.",
            f"Study practical applications of {topic}.",
            f"Analyze relevant data related to {topic}.",
            f"Explore future improvements in {topic}."
        ]
    }

    string_keys = [
        "title",
        "summary",
        "problem",
        "methodology",
        "expected_outcomes"
    ]

    list_keys = ["objectives", "research_ideas"]

    for key in string_keys:
        value = data.get(key)

        # Convert a list of outcome sentences into one string.
        if key == "expected_outcomes" and isinstance(value, list):
            value = " ".join(
                str(item).strip()
                for item in value
                if isinstance(item, str) and item.strip()
            )

        if not isinstance(value, str) or not value.strip():
            data[key] = defaults[key]
        else:
            data[key] = value.strip()

    for key in list_keys:
        value = data.get(key)

        if isinstance(value, str):
            value = [value]

        if not isinstance(value, list):
            data[key] = defaults[key]
            continue

        cleaned = [
            item.strip()
            for item in value
            if isinstance(item, str) and item.strip()
        ]

        data[key] = cleaned if cleaned else defaults[key]

    return data


def read_form():
    topic = request.form.get("topic", "").strip()
    field = request.form.get("field", "").strip()
    level = request.form.get("level", "").strip()
    description = request.form.get("description", "").strip()

    if not topic:
        raise ValueError("Please enter a research topic.")

    if len(topic) > 200:
        raise ValueError("Research topic must be 200 characters or fewer.")

    if len(field) > 150:
        raise ValueError("Research field must be 150 characters or fewer.")

    if len(level) > 100:
        raise ValueError("Academic level must be 100 characters or fewer.")

    if len(description) > 1000:
        raise ValueError(
            "Additional requirements must be 1000 characters or fewer."
        )

    return topic, field, level, description


def generate_ai_proposal(topic, field, level, description):
    prompt = f"""
Create a concise academic research proposal using these details.

Topic: {topic}
Field: {field or "General Studies"}
Academic level: {level or "Not specified"}
Requirements: {description or "None"}

Return one valid JSON object with exactly these seven keys:
- title: a short string
- summary: a short string
- problem: a short problem statement string
- objectives: an array of exactly three short strings
- methodology: a concise string describing the research steps
- expected_outcomes: a short string, NOT an array
- research_ideas: an array of three short strings

Use information relevant to the topic.
Do not use placeholder text.
Do not duplicate keys.
Do not add Markdown or text outside the JSON object.
Keep the response short.
"""

    try:
        response = requests.post(
            OLLAMA_URL,
            json={
                "model": MODEL_NAME,
                "prompt": prompt,
                "format": "json",
                "stream": False,
                "options": {
                    "temperature": 0.1,
                    "num_predict": 400,
                    "repeat_last_n": 32,
                    "repeat_penalty": 1.2
                }
            },
            timeout=OLLAMA_TIMEOUT
        )

        if not response.ok:
            app.logger.error(
                "Ollama HTTP status: %s", response.status_code
            )
            app.logger.error(
                "Ollama error response: %s", response.text
            )
            raise RuntimeError(
                f"Ollama returned error {response.status_code}. "
                "Please try again."
            )

    except requests.exceptions.RequestException as error:
        app.logger.exception("Ollama request failed.")
        raise RuntimeError(
            "Could not connect to Ollama. Make sure Ollama is running."
        ) from error

    try:
        result = response.json()
        raw_response = result.get("response", "")

        if not raw_response.strip():
            raise ValueError("Ollama returned an empty response.")

        proposal = json.loads(raw_response)
        proposal = validate_proposal(proposal, topic)

        return proposal

    except (ValueError, TypeError, AttributeError) as error:
        app.logger.exception("Could not read the AI response.")
        raise ValueError(
            "The AI response could not be processed. Please try again."
        ) from error


def create_word_document(topic, field, level, description, proposal):
    doc = Document()

    doc.add_heading("AI Research Proposal", 0)

    doc.add_heading("Research Title", level=1)
    doc.add_paragraph(proposal["title"])

    doc.add_heading("Research Topic", level=1)
    doc.add_paragraph(topic)

    doc.add_heading("Research Field", level=1)
    doc.add_paragraph(field or "Not specified")

    doc.add_heading("Academic Level", level=1)
    doc.add_paragraph(level or "Not specified")

    if description:
        doc.add_heading("Additional Requirements", level=1)
        doc.add_paragraph(description)

    sections = [
        ("Research Summary", "summary"),
        ("Problem Statement", "problem"),
        ("Methodology", "methodology"),
        ("Expected Outcomes", "expected_outcomes")
    ]

    for heading, key in sections:
        doc.add_heading(heading, level=1)
        doc.add_paragraph(proposal[key])

    doc.add_heading("Research Objectives", level=1)

    for objective in proposal["objectives"]:
        doc.add_paragraph(objective, style="List Bullet")

    doc.add_heading("Suggested Research Ideas", level=1)

    for idea in proposal["research_ideas"]:
        doc.add_paragraph(idea, style="List Number")

    output = BytesIO()
    doc.save(output)
    output.seek(0)

    return output


@app.route("/")
def home():
    return render_template("index.html")


@app.route("/generate", methods=["POST"])
def generate():
    try:
        topic, field, level, description = read_form()

        proposal = generate_ai_proposal(
            topic, field, level, description
        )

        return render_template(
            "result.html",
            topic=topic,
            field=field,
            level=level,
            description=description,
            proposal=proposal,
            research_ideas=proposal["research_ideas"]
        )

    except (ValueError, RuntimeError) as error:
        app.logger.error("Proposal generation failed: %s", error)
        return str(error), 502

    except Exception:
        app.logger.exception("Unexpected generation error.")
        return (
            "Unexpected error. Check the VS Code terminal.",
            500
        )


@app.route("/download", methods=["POST"])
def download_proposal():
    try:
        topic, field, level, description = read_form()

        # Use the existing proposal. Do not call Ollama again.
        raw = request.form.get("proposal_json", "")

        if not raw or len(raw) > 50000:
            return (
                "Proposal data is missing. Please generate it again.",
                400
            )

        proposal = json.loads(raw)
        proposal = validate_proposal(proposal, topic)

        output = create_word_document(
            topic, field, level, description, proposal
        )

        return send_file(
            output,
            as_attachment=True,
            download_name="research_proposal.docx",
            mimetype=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            )
        )

    except (ValueError, TypeError, json.JSONDecodeError) as error:
        app.logger.exception("Invalid proposal for download.")
        return f"Could not process proposal: {error}", 400

    except Exception:
        app.logger.exception("Word download failed.")
        return (
            "Could not create the Word document. "
            "Check the VS Code terminal.",
            500
        )


if __name__ == "__main__":
    app.run(debug=True)
